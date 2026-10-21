"""Excel parsing and export with openpyxl."""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .masters import MASTERS, Column, Master, norm_header, slug


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _find_header(grid: list[list[str]], columns: list[Column]):
    """Return (header_row_index, {col_key: col_index}) or (None, positional map)."""
    need = min(2, len(columns))
    for r, row in enumerate(grid[:10]):
        heads = [norm_header(v) for v in row]
        mapping, used = {}, set()
        for c in columns:
            names = {norm_header(a) for a in c.aliases}
            for i, h in enumerate(heads):
                if h and i not in used and h in names:
                    mapping[c.key] = i
                    used.add(i)
                    break
        if len(mapping) >= need:
            # Columns without a recognised header keep their position if that cell is free.
            for pos, c in enumerate(columns):
                if c.key not in mapping and pos not in used:
                    mapping[c.key] = pos
            return r, mapping
    return None, {c.key: i for i, c in enumerate(columns)}


def _rows_from_grid(grid: list[list[str]], master: Master, columns: list[Column] | None):
    """Return (columns, rows) where rows are dicts keyed by column key."""
    grid = [row for row in grid if any(row)]
    if not grid:
        return columns or [], []

    if columns is None:  # dynamic columns: first non-empty row is the header
        taken: set[str] = set()
        header = grid[0]
        width = max(i for i, v in enumerate(header) if v) + 1
        columns = [Column(slug(h or f"Column {i + 1}", taken), h or f"Column {i + 1}") for i, h in enumerate(header[:width])]
        start, mapping = 1, {c.key: i for i, c in enumerate(columns)}
    else:
        hdr, mapping = _find_header(grid, columns)
        start = 0 if hdr is None else hdr + 1

    rows = []
    for row in grid[start:]:
        rec = {c.key: (row[mapping[c.key]] if mapping[c.key] < len(row) else "") for c in columns}
        if master.keyed and not rec[columns[0].key]:
            continue  # keyed masters need a code
        if any(rec.values()):
            rows.append(rec)
    return columns, rows


EXCEL_ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def _read_sheet(ws) -> list[list[str]]:
    """Master sheets as text. Formula errors such as #N/A are read as blank cells."""
    return [["" if t in EXCEL_ERRORS else t for t in map(cell_text, row)] for row in ws.iter_rows(values_only=True)]


def _open(data: bytes):
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # not a valid xlsx
        raise ValueError(f"Could not read the Excel file ({e.__class__.__name__}). Please upload a .xlsx file.") from e


def parse_for_type(data: bytes, master: Master):
    """Single-master upload: use the matching tab, else the first sheet."""
    wb = _open(data)
    try:
        name = next((n for n in wb.sheetnames if master.sheet.search(n)), wb.sheetnames[0])
        cols, rows = _rows_from_grid(_read_sheet(wb[name]), master, master.columns)
        return name.strip(), cols, rows
    finally:
        wb.close()


def parse_workbook(data: bytes):
    """Full workbook upload: every master whose tab is found. Returns {type: (sheet, columns, rows)}."""
    wb = _open(data)
    try:
        out = {}
        for t, m in MASTERS.items():
            name = next((n for n in wb.sheetnames if m.sheet.search(n)), None)
            if name:
                cols, rows = _rows_from_grid(_read_sheet(wb[name]), m, m.columns)
                out[t] = (name.strip(), cols, rows)
        return out
    finally:
        wb.close()


def build_xlsx(master: Master, columns: list[Column], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = master.label[:31].replace("&", "and")
    ws.append([c.label for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE6F3")
    for r in rows:
        # Write plain digit strings back as numbers so Excel lookups keep working.
        ws.append([int(v) if isinstance(v, str) and v.isdigit() and len(v) < 16 else v for v in (r.get(c.key, "") for c in columns)])
    for i, c in enumerate(columns):
        width = max([len(c.label)] + [len(str(r.get(c.key, ""))) for r in rows[:2000]])
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = min(60, max(10, width + 2))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
