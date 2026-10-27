"""Full workbook of a run: one Excel file with every tab the work needs.

  REDEFINED DATA   the mapped result (all columns, coloured flags, filter)
  SAP DATA         the SAP sheet exactly as it was uploaded
  <masters>        every master as it was when the file was processed (descriptions, D1, F1)
  Summary          the run summary

The masters are a snapshot, so the file still explains the result after the masters change.
"""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..db import master_repo
from ..domain.masters import MASTERS, Master
from . import sap

_BAD = re.compile(r"[\[\]:*?/\\]")
HEAD_FONT, HEAD_FILL = Font(bold=True), PatternFill("solid", fgColor="DDE6F3")


def _tab_name(label: str, used: set[str]) -> str:
    """Excel tab names: at most 31 characters, none of []:*?/\\, unique ignoring case."""
    base = _BAD.sub("-", label).strip()[:31].strip() or "Sheet"
    name, i = base, 2
    while name.lower() in used:
        suffix = f" {i}"
        name, i = base[:31 - len(suffix)].strip() + suffix, i + 1
    used.add(name.lower())
    return name


def _head(ws, labels):
    cells = []
    for h in labels:
        c = WriteOnlyCell(ws, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        cells.append(c)
    ws.append(cells)


def _widths(ws, labels, rows):
    for i, h in enumerate(labels):
        longest = max([len(str(h))] + [len(str(r[i])) for r in rows[:500] if i < len(r) and r[i] is not None])
        ws.column_dimensions[get_column_letter(i + 1)].width = min(60, max(10, longest + 2))


def _write_sap_data(wb, raw: dict, used: set[str]):
    ws = wb.create_sheet(_tab_name("SAP DATA", used))
    ws.freeze_panes = "A2"
    header, body = raw["header"], raw["body"]
    _widths(ws, header, body)
    _head(ws, header)
    for r in body:
        row = []
        for v in r:
            if isinstance(v, (datetime, date)):  # dates as in the upload, without a 0:00:00 time part
                c = WriteOnlyCell(ws, value=v)
                c.number_format = "dd-mm-yyyy" if not isinstance(v, datetime) or v.time() == datetime.min.time() else "dd-mm-yyyy hh:mm"
                row.append(c)
            else:
                row.append(v)
        ws.append(row)


def _write_master(wb, m: Master, used: set[str]):
    cols, rows = master_repo.columns_for(m), master_repo.all_rows(m)
    if not cols or not rows:  # an empty master (for example D1 ESI & PF) has nothing to show
        return
    ws = wb.create_sheet(_tab_name(m.label, used))
    ws.freeze_panes = "A2"
    table = [[r.get(c.key, "") for c in cols] for r in rows]
    labels = [c.label for c in cols]
    _widths(ws, labels, table)
    _head(ws, labels)
    for r in table:  # plain digit strings go back as numbers so Excel lookups keep working
        ws.append([int(v) if isinstance(v, str) and v.isdigit() and len(v) < 16 else v for v in r])


def build_full_output(summary, out_header, out_rows, raw: dict) -> bytes:
    wb = Workbook(write_only=True)
    used = {"redefined data", "summary"}
    sap._write_redefined(wb, out_header, out_rows)
    _write_sap_data(wb, raw, used)
    for m in MASTERS.values():
        _write_master(wb, m, used)
    sap._write_summary(wb, summary)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
