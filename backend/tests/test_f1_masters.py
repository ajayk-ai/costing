"""F1 masters (GL Nature, Overhead): loaded from the workbook tabs and listed under their own group."""
import io

from openpyxl import Workbook, load_workbook

from conftest import make_xlsx, xlsx_file


def workbook(gl_rows, oh_rows):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("F1 GL NATURE")  # tab names as in the expense report, with stray spaces
    ws.append(["Account", "Account Description", "Nature", "GL nature"])
    for r in gl_rows:
        ws.append(r)
    ws = wb.create_sheet(" F1 OVERHEAD ")
    ws.append(["GL + Department", "Overhead"])
    for r in oh_rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": ("f1.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


GL = [[39500100.0, "Gain from exchange rate fluctuation", "Fixed", "Finance Charge"],
      [42201200.0, "Overtime Paid to Workmen", "Variable", "Labour Expenses"]]
OH = [["Audit feesFinance", "Admin OH"],
      ["Assembly ChargesProduction", "Direct labour"],
      ["Audit feesFinance", "Admin OH"]]  # exact duplicate key


def test_f1_masters_are_listed_in_their_own_group(client):
    ms = {m["type"]: m for m in client.get("/api/masters").json()}
    assert ms["f1_gl_nature"]["group"] == ms["f1_overhead"]["group"] == "F1 Mapping"
    assert [c["label"] for c in ms["f1_gl_nature"]["columns"]] == ["Account", "Account Description", "Nature", "GL nature"]
    assert [c["label"] for c in ms["f1_overhead"]["columns"]] == ["GL + Department", "Overhead"]


def test_workbook_upload_loads_both_f1_tabs(client):
    r = client.post("/api/upload-workbook", files=workbook(GL, OH))
    assert r.status_code == 200, r.text
    res = r.json()["results"]
    assert res["f1_gl_nature"]["inserted"] == 2 and res["f1_overhead"]["inserted"] == 2
    assert res["f1_overhead"]["skipped"] == 1  # the duplicate key
    rows = client.get("/api/masters/f1_gl_nature?q=39500100").json()["rows"]
    assert len(rows) == 1 and rows[0]["account"] == "39500100"  # 39500100.0 becomes a clean code
    assert rows[0]["gl_nature"] == "Finance Charge" and rows[0]["nature"] == "Fixed"
    oh = client.get("/api/masters/f1_overhead?q=Audit").json()["rows"]
    assert [(x["gl_department"], x["overhead"]) for x in oh] == [("Audit feesFinance", "Admin OH")]


def test_merge_updates_changed_values_and_blank_never_overwrites(client):
    gl = [[39500100.0, "Gain from exchange rate fluctuation", "Variable", ""]]  # Nature changed, GL nature blank
    r = client.post("/api/upload-workbook", files=workbook(gl, [["Audit feesFinance", "Selling OH"]]))
    assert r.status_code == 200, r.text
    row = client.get("/api/masters/f1_gl_nature?q=39500100").json()["rows"][0]
    assert row["nature"] == "Variable" and row["gl_nature"] == "Finance Charge"
    oh = client.get("/api/masters/f1_overhead?q=Audit").json()["rows"][0]
    assert oh["overhead"] == "Selling OH"


def test_f1_masters_export_and_template(client):
    for t in ("f1_gl_nature", "f1_overhead"):
        assert client.get(f"/api/masters/{t}/export").status_code == 200
        assert client.get(f"/api/masters/{t}/template").status_code == 200


def test_sap_result_gets_gl_nature_and_overhead(client):
    """GL Nature by Account; Overhead by F1 account description + Department; unmatched stay blank and are listed."""
    # 400200 -> Logistics is also loaded by the end-to-end tests; 409 just means it is already there
    client.post("/api/masters/d1_fixed", json={"account": "400200", "account_desc": "Freight", "department": "Logistics"})
    for t, row in (
        ("f1_gl_nature", {"account": "42022100", "account_desc": "Security fees", "nature": "Fixed", "gl_nature": "Administration Expenses"}),
        ("f1_overhead", {"gl_department": "Security feesFinance", "overhead": "Admin OH"}),
        ("d1_fixed", {"account": "42022100", "account_desc": "Security fees", "department": "Finance"}),
        ("d1_fixed", {"account": "42201200", "account_desc": "Overtime", "department": "Production"}),
    ):
        assert client.post(f"/api/masters/{t}", json=row).status_code == 201
    rows = [
        ["", "42022100", "", "", "", "", "2026/01", "10", "audit"],     # Finance -> key "Security feesFinance" -> Admin OH
        ["", "42201200", "", "", "", "", "2026/01", "20", "overtime"],  # key "Overtime Paid to WorkmenProduction" not in F1 Overhead
        ["", "400200", "", "", "", "", "2026/01", "30", "freight"],     # Logistics, but the account is not in F1 GL Nature
        ["", "999999", "", "", "", "", "2026/01", "40", "unmapped"],    # no department at all
    ]
    sap = make_xlsx({"SAP DATA": [["Offst.acct", "Account", "Cost Ctr", "Order", "Material", "Assignment",
                                   "Year/month", "Amount in local cur.", "Text"], *rows]})
    r = client.post("/api/sap/process", files=xlsx_file(sap, "f1.xlsx"))
    assert r.status_code == 200, r.text
    run = r.json()
    h = run["summary"]["header"]
    d = h.index("Department")
    assert h[d:d + 3] == ["Department", "GL Nature", "Overhead"]

    got = {x[h.index("Account")]: (x[h.index("GL Nature")], x[h.index("Overhead")]) for x in run["preview"]["all"]}
    assert got["42022100"] == ("Administration Expenses", "Admin OH")
    assert got["42201200"] == ("Labour Expenses", "")  # GL nature found, overhead combination missing
    assert got["400200"] == ("", "") and got["999999"] == ("", "")

    f1 = run["summary"]["f1"]
    assert f1["glNature"]["matched"] == 2 and f1["glNature"]["missingRows"] == 2 and f1["glNature"]["blank"] == 0  # accounts 400200 and 999999
    assert f1["overhead"]["matched"] == 1 and f1["overhead"]["noDepartment"] == 1
    assert f1["overhead"]["noGlAccount"] == 1 and f1["overhead"]["missingRows"] == 1
    miss = f1["overhead"]["missing"][0]
    assert (miss["description"], miss["department"]) == ("Overtime Paid to Workmen", "Production")
    assert miss["key"] == "Overtime Paid to WorkmenProduction" and miss["rows"] == 1

    ws = load_workbook(io.BytesIO(client.get(f"/api/sap/runs/{run['id']}/download").content))
    sheet = ws.worksheets[0]
    head = [c.value for c in sheet[1]]
    assert head[d + 1:d + 3] == ["GL Nature", "Overhead"]
    assert [sheet.cell(2, d + 2).value, sheet.cell(2, d + 3).value] == ["Administration Expenses", "Admin OH"]
    texts = [c.value for row in ws["Summary"].iter_rows() for c in row[:1]]
    assert "F1 mapping" in texts and "Add to F1 Overhead: GL description" in texts


def _process(client, rows, name="full.xlsx"):
    sap = make_xlsx({"SAP DATA": [["Offst.acct", "Account", "Cost Ctr", "Order", "Material", "Assignment",
                                   "Year/month", "Amount in local cur.", "Text"], *rows]})
    r = client.post("/api/sap/process", files=xlsx_file(sap, name))
    assert r.status_code == 200, r.text
    return r.json()


ROWS = [["", "42022100", "", "", "", "", "2026/01", "10", "audit"],
        ["", "400200", "", "", "", "", "2026/01", "30", "freight"]]


def _ensure(client, master, row):
    """Add a master row if it is not there yet (409 = already loaded by an earlier test)."""
    assert client.post(f"/api/masters/{master}", json=row).status_code in (201, 409)


def test_full_workbook_has_all_tabs(client):
    _ensure(client, "material", {"code": "0000100", "description": "Bolt M8"})
    _ensure(client, "account", {"code": "400100", "description": "Salary"})
    _ensure(client, "d1_fixed", {"account": "400200", "account_desc": "Freight", "department": "Logistics"})
    _ensure(client, "d1_emp_id", {"emp_id": "E77", "revised_dept": "Assembly"})
    run = _process(client, ROWS)
    assert run["hasFull"] is True
    dl = client.get(f"/api/sap/runs/{run['id']}/download-full")
    assert dl.status_code == 200 and dl.headers["content-disposition"].endswith('_FULL.xlsx"')
    wb = load_workbook(io.BytesIO(dl.content))
    names = wb.sheetnames
    assert names[:2] == ["REDEFINED DATA", "SAP DATA"] and names[-1] == "Summary"
    for tab in ("Material Description", "Account Description", "D1 Fixed Expenses based Account", "D1 Emp ID",
                "F1 GL Nature", "F1 Overhead"):
        assert tab in names
    assert "D1 ESI & PF" not in names  # an empty master has no tab
    assert all(len(n) <= 31 for n in names) and len({n.lower() for n in names}) == len(names)

    sap = wb["SAP DATA"]  # exactly the uploaded sheet
    assert [c.value for c in sap[1]][:3] == ["Offst.acct", "Account", "Cost Ctr"] and sap.max_row == len(ROWS) + 1
    assert [sap.cell(i, 2).value for i in (2, 3)] == ["42022100", "400200"] or [sap.cell(i, 2).value for i in (2, 3)] == [42022100, 400200]
    redef = wb["REDEFINED DATA"]
    head = [c.value for c in redef[1]]
    assert head[-3:] == ["Department", "GL Nature", "Overhead"] and redef.max_row == len(ROWS) + 1  # no AI model in tests
    gl = wb["F1 GL Nature"]
    assert [c.value for c in gl[1]] == ["Account", "Account Description", "Nature", "GL nature"] and gl.max_row > 2
    assert [c.value for c in wb["F1 Overhead"][1]] == ["GL + Department", "Overhead"]
    # the plain result file is unchanged: two tabs only
    plain = load_workbook(io.BytesIO(client.get(f"/api/sap/runs/{run['id']}/download").content))
    assert plain.sheetnames == ["REDEFINED DATA", "Summary"]


def test_full_workbook_is_offered_only_when_the_file_exists(client):
    from app.services import sap_runs

    run = _process(client, ROWS)
    listed = {x["id"]: x for x in client.get("/api/sap/runs").json()}
    assert listed[run["id"]]["hasFull"] is True
    sap_runs.full_path(run).unlink()  # like a run made before the full workbook existed
    assert client.get(f"/api/sap/runs/{run['id']}").json()["hasFull"] is False
    assert {x["id"]: x for x in client.get("/api/sap/runs").json()}[run["id"]]["hasFull"] is False
    r = client.get(f"/api/sap/runs/{run['id']}/download-full")
    assert r.status_code == 404 and "Process the SAP file again" in r.json()["detail"]
    assert client.get(f"/api/sap/runs/{run['id']}/download").status_code == 200  # the result file still works
    assert client.get("/api/sap/runs/999999/download-full").status_code == 404


def test_deleting_a_run_deletes_both_files(client):
    from app.services import sap_runs

    run = _process(client, ROWS)
    result, full = sap_runs.output_path(run), sap_runs.full_path(run)
    assert result and full
    assert client.delete(f"/api/sap/runs/{run['id']}").status_code == 200
    assert not result.exists() and not full.exists()
